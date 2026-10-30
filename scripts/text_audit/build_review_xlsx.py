# สร้างไฟล์ Excel ให้เจ้าของภาษา/ครูตรวจเนื้อหา: node scripts/text_audit/extract_review.mjs rows.json && python3 scripts/text_audit/build_review_xlsx.py rows.json out.xlsx
import json, re, sys
import os
HERE=os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0,HERE)
from g2p import g2p
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

rows=json.load(open(sys.argv[1]))  # จาก: node scripts/text_audit/extract_review.mjs rows.json
F='Arial'
def font(**k): return Font(name=F,size=k.pop('size',10),**k)
HEAD=PatternFill('solid',fgColor='1F3864'); INPUT=PatternFill('solid',fgColor='FFF9C4'); INHEAD=PatternFill('solid',fgColor='BF9000'); GREY=PatternFill('solid',fgColor='F2F2F2')
thin=Side(style='thin',color='BFBFBF'); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='top')

PRIORITY={'บทเรียน':'สูง','เรื่องราว':'สูง','กิจวัตร':'กลาง','สัตว์เลี้ยง':'กลาง','ภารกิจ/ความสำเร็จ':'กลาง','ร้านค้า':'ต่ำ'}
SENT_END=set('요다까죠네니세습십군걸데지')
def flags(r):
    out=[]
    ko=r['ko']; hang=[c for c in re.sub(r'[^가-힣]','',ko)]
    pr=g2p(ko) if hang else ''
    spoken = pr if pr and pr!=ko.replace('?','').replace('!','') else ''
    is_sentence = ('เรื่องราว' == r['sheet'] and 'คำศัพท์' not in r['where']) or '· ประโยค' in r['where']
    if is_sentence and hang and hang[-1] not in SENT_END and len(hang)>=3:
        out.append('ลงท้ายไม่เป็นแบบสุภาพ (ภาษาพูดกันเอง?) — โปรดยืนยันว่าเหมาะกับผู้พูดคนนี้')
    if r['note']: out.append(r['note'])
    return spoken,' / '.join(out)

COLS=[('#',5),('ความสำคัญ',9),('ที่มา (หน้า/บท/ฉาก)',34),('ผู้พูด',11),('한국어 (เกาหลี)',34),('เสียงอ่านจริง (สร้างอัตโนมัติ — กฎการออกเสียง)',24),('คำอ่านไทย (ที่แอปแสดง)',30),('คำแปลไทย (ที่แอปแสดง)',34),('หมายเหตุจากผู้พัฒนา',30)]
INCOLS=[('① 한국어 ถูกต้อง/เป็นธรรมชาติ?',16),('① ข้อเสนอแก้ 한국어',26),('② คำอ่านไทยตรงกับเสียงจริง?',16),('② ข้อเสนอคำอ่านไทย',24),('③ คำแปลไทยถูกต้อง?',16),('③ ข้อเสนอคำแปล',24),('ความเห็นอื่น ๆ',28)]
dv_ko=DataValidation(type='list',formula1='"✔ ถูกต้อง,✎ ควรแก้,? ไม่แน่ใจ"',allow_blank=True)

wb=Workbook()
# ---- guide
g=wb.active; g.title='อ่านก่อน'
g.column_dimensions['A'].width=4; g.column_dimensions['B'].width=34; g.column_dimensions['C'].width=96
def put(r,c,v,**k):
    cell=g.cell(r,c,v); cell.font=font(**k); cell.alignment=Alignment(wrap_text=True,vertical='top'); return cell
put(1,2,'รายการตรวจเนื้อหาแอป 형범 (Hanguk Hyeong-Bum)',bold=True,size=16)
put(2,2,'แอปเรียนภาษาเกาหลีสำหรับคนไทย — ไฟล์นี้รวมข้อความเกาหลี คำอ่านไทย และคำแปลทั้งแอป เพื่อให้เจ้าของภาษา/ครูช่วยตรวจก่อนเผยแพร่กว้างขึ้น',size=10)
r=4
sections=[
('ใครตรวจอะไร',[
 ('ผู้ตรวจที่เป็นคนเกาหลี (또는 한국어 선생님)','ตรวจช่อง ①: ภาษาเกาหลีถูกต้อง/เป็นธรรมชาติ/ระดับความสุภาพเหมาะกับผู้พูดหรือไม่ · 한국어 칸 ①만 확인해 주세요 (아래 한국어 안내 참고)'),
 ('ผู้ตรวจที่อ่านไทยออก (คนไทยที่รู้เกาหลี หรือครูสอนเกาหลี)','ตรวจช่อง ② คำอ่านไทยว่าตรงกับ "เสียงที่ได้ยินจริง" และช่อง ③ คำแปลไทยว่าตรงความหมายและเป็นภาษาไทยที่เป็นธรรมชาติ'),
 ('ถ้าตรวจคนเดียวทั้งสามช่อง','ได้เลย — แต่ถ้าไม่แน่ใจข้อไหน เลือก "? ไม่แน่ใจ" แล้วเขียนเหตุผล ดีกว่าเดาให้'),
]),
('วิธีกรอก',[
 ('ช่องสีเหลือง','คือช่องที่ต้องกรอก ช่องสีเทา/ขาวเป็นข้อมูลของแอป ห้ามแก้ (แก้ในช่อง "ข้อเสนอ" แทน)'),
 ('ตัวเลือกใน ① ② ③','✔ ถูกต้อง · ✎ ควรแก้ (ใส่ข้อเสนอในช่องถัดไป) · ? ไม่แน่ใจ — เว้นว่างได้ถ้ายังไม่ได้ตรวจแถวนั้น'),
 ('ลำดับความสำคัญ','ตรวจแถวที่ "สูง" ก่อน (บทเรียนและเรื่องราว ผู้เรียนเห็นบ่อยที่สุด) แล้วค่อย "กลาง" และ "ต่ำ" (ชื่อไอเทมในร้านค้า) — ใช้ตัวกรองที่หัวคอลัมน์ได้'),
 ('แถวที่มีหมายเหตุ','ผู้พัฒนาเขียนเตือนไว้ เช่น ประโยคที่ลงท้ายไม่สุภาพ (ตั้งใจให้เป็นคำพูดกันเอง?) หรือคำที่เสียงเปลี่ยนจากตัวเขียน'),
]),
('กฎการถอดเสียงอ่านเกาหลีเป็นไทยที่แอปใช้ (ช่อง ②)',[
 ('พยัญชนะต้น (เสียงเบา)','ㄱ=ก · ㄷ=ด · ㅂ=บ · ㅈ=จ'),
 ('พยัญชนะต้น (พ่นลม)','ㅋ=ค · ㅌ=ท · ㅍ=พ · ㅊ=ช'),
 ('พยัญชนะต้น (เสียงตึง)','ㄲ=ก · ㄸ=ต · ㅃ=ป · ㅉ=จ  (ภาษาไทยแยกตึง/เบาไม่ได้ จึงให้เสียงเป็นตัวสนับสนุน)'),
 ('สระ','ㅐ=แ · ㅔ=เ · ㅚ=เว · ㅙ=แว · ㅓ=ออ · ㅗ=โอ · ㅡ=อือ · สระนำ (เ แ โ ไ) เขียนหน้าพยัญชนะเสมอ'),
 ('เขียนตามเสียงจริง','เช่น 많아요 → มานาโย · 먹었어요 → มอกอซซอโย (ไม่ใช่ตามตัวเขียน) — ช่อง "เสียงอ่านจริง" ที่สร้างอัตโนมัติไว้ช่วยเทียบ (ผ่านกฎการออกเสียงหลัก อาจผิดในบางคำ)'),
 ('สิ่งที่ภาษาไทยทำไม่ได้','ไม่มีเสียงสั้น-ยาวของสระ ไม่มี ㅡ/ㅓ ที่ตรงเป๊ะ และเสียงวรรณยุกต์ไม่ตรงกับน้ำเสียงเกาหลี — คำอ่านไทยเป็นแค่ตัวช่วยเริ่มต้น ผู้เรียนมีเสียงพากย์ฟังคู่กัน'),
]),
('ขอบเขต / สิ่งที่ไม่ต้องตรวจ',[
 ('ไม่ต้องตรวจ','ภาพ เกม ราคา ระบบรางวัล · ชื่อไอเทมที่เป็นชื่อเฉพาะทำนองสมมติ (ตรวจแค่ว่าเกาหลีใช้พูดกันจริงไหม)'),
 ('ภาษาเกาหลีมาตรฐานที่ใช้','เกาหลีใต้ (서울말) สุภาพแบบ -아요/어요 เป็นหลัก ใช้ -ㅂ니다 ในประโยคทางการ (ทักทาย/ขอบคุณ)'),
 ('ผู้พูดในเรื่อง','คุณยาย (할머니) พูดกับเด็กอย่างอบอุ่น · แม่ค้า · ผู้เล่น (นักเรียนต่างชาติ) · ผู้เล่าเรื่อง · สัตว์เลี้ยง (ใช้ภาษาเด็ก/น่ารักได้) — ดูช่อง "ผู้พูด"'),
]),
('한국어 안내 (Korean reviewers)',[
 ('이 파일은 무엇인가요?','태국인을 위한 한국어 학습 앱의 모든 한국어 문장과 단어입니다. 원어민의 시각에서 어색하거나 틀린 표현이 없는지 확인해 주세요.'),
 ('확인해 주실 부분','노란색 칸 ① 「한국어 자연스러움」만 확인해 주시면 됩니다. 태국어(번역·발음 표기) 칸은 건너뛰셔도 됩니다.'),
 ('표시 방법','✔ 올바름 / ✎ 수정 필요 (옆 칸에 고친 문장을 써 주세요) / ? 잘 모르겠음'),
 ('참고','화자(할머니·가게 주인·학습자·내레이터·반려동물)에 맞는 말투인지, 띄어쓰기·철자·존댓말 수준이 일관적인지, 실제로 그렇게 말하는지 봐 주세요. 초급 학습자를 위한 문장이라 아주 단순하게 쓴 점은 양해 부탁드립니다.'),
]),
]
for title,items in sections:
    c=put(r,2,title,bold=True,size=12); c.fill=GREY; g.cell(r,3).fill=GREY; r+=1
    for a,b in items:
        put(r,2,a,bold=True); put(r,3,b); g.row_dimensions[r].height=max(18,15*(1+len(b)//95)); r+=1
    r+=1
put(r,2,'ตัวอย่างการกรอก (แถวตัวอย่าง ไม่ใช่ข้อมูลจริง)',bold=True)
r+=1
for i,v in enumerate(['① ✎ ควรแก้','ตัวอย่าง: 안녕히 계세요 (ถ้าผู้พูดเป็นคนที่จากไป)','② ✔ ถูกต้อง','','③ ? ไม่แน่ใจ','คำแปลฟังดูเป็นทางการเกินไปสำหรับเด็ก']):
    pass
put(r,2,'ช่อง ① = ✎ ควรแก้'); put(r,3,'ข้อเสนอ 한국어: 안녕히 가세요 (ผู้พูดเป็นคนที่ "อยู่" จึงใช้ 가세요 กับคนที่กำลังไป)'); r+=1
put(r,2,'ช่อง ③ = ? ไม่แน่ใจ'); put(r,3,'ความเห็นอื่น ๆ: คำแปล "ลาก่อน" ฟังทางการเกินไปสำหรับบทสนทนาของเด็ก'); r+=1

# ---- review sheets
order=['บทเรียน','เรื่องราว','กิจวัตร','สัตว์เลี้ยง','ภารกิจ/ความสำเร็จ','ร้านค้า']
summary=[]
for sh in order:
    ws=wb.create_sheet(sh.replace('/','-'))
    rs=[x for x in rows if x['sheet']==sh]
    heads=[c for c,_ in COLS]+[c for c,_ in INCOLS]
    for j,(h,w) in enumerate(COLS+INCOLS,1):
        cell=ws.cell(1,j,h); cell.font=font(bold=True,color='FFFFFF'); cell.fill=INHEAD if j>len(COLS) else HEAD
        cell.alignment=Alignment(wrap_text=True,vertical='center'); cell.border=BORDER; ws.column_dimensions[L(j)].width=w
    ws.row_dimensions[1].height=48
    for i,x in enumerate(rs,1):
        spoken,note=flags(x)
        vals=[i,PRIORITY[sh],x['where'],x['who'],x['ko'],spoken,x['read'],x['th'],note]
        for j,v in enumerate(vals,1):
            c=ws.cell(i+1,j,v); c.font=font(size=11 if j==5 else 10); c.alignment=wrap; c.border=BORDER
        for j in range(len(COLS)+1,len(COLS)+len(INCOLS)+1):
            c=ws.cell(i+1,j); c.fill=INPUT; c.border=BORDER; c.font=font(); c.alignment=wrap
    n=len(rs)
    for col in (len(COLS)+1,len(COLS)+3,len(COLS)+5):
        dv=DataValidation(type='list',formula1='"✔ ถูกต้อง,✎ ควรแก้,? ไม่แน่ใจ"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'{L(col)}2:{L(col)}{n+1}')
    ws.freeze_panes='F2'; ws.auto_filter.ref=f'A1:{L(len(COLS)+len(INCOLS))}{n+1}'
    summary.append((ws.title,n))

# ---- listening sheet
ls=wb.create_sheet('ฟังเสียง')
heads=['เสียงพูดที่อยากให้ฟัง (เปิดแอป → ปุ่ม 🔊 ของคำนั้น)','ข้อความ','วิธีไปฟัง','เหตุผลที่ขอให้ฟัง','ฟังแล้วเป็นอย่างไร?','ความเห็น']
for j,(h,w) in enumerate(zip(heads,[24,18,40,44,18,36]),1):
    c=ls.cell(1,j,h); c.font=font(bold=True,color='FFFFFF'); c.fill=INHEAD if j>4 else HEAD; c.alignment=Alignment(wrap_text=True,vertical='center'); ls.column_dimensions[L(j)].width=w
ls.row_dimensions[1].height=40
listen=[('오래','เรื่องตอน 1 · ช่วง 알 (คุณยายพูดว่า "이 알은 오래 기다렸어요") → หน้าสรุปคำศัพท์ ปุ่ม 🔊','เครื่องตรวจอัตโนมัติจับคำได้ต่ำ (ครูมินจุน 0%, ครูฮานา 50%) อาจเพี้ยนจริงหรืออาจเป็นข้อจำกัดของเครื่องตรวจ'),
('곶감','เรื่องตอน 1 · ช่วง "นิทานเสือกับลูกพลับแห้ง" (c1_tale) และคำตอบในตอน 3','คำที่เจอยากสำหรับเครื่อง (ทั้งสองเสียง 50%)'),
('익어요','เรื่องตอน 3 · ช่วงทำซงพยอน ("솥에서 송편이 익어 가요")','ทั้งสองเสียง 33%'),
('어떻게','เรื่องตอน 3 · ช่วงทำซงพยอน ("어떻게 접었어요?")','ครูฮานา 33%'),
('돼요','เรื่องตอน 1 ("닭이 돼요") และตอน 3 ("이렇게 하면 돼요?") · บทเรียนร้านค้า ("카드 돼요?")','ครูฮานา 50%'),
('옛날 옛적에','เรื่องตอน 1 · นิทานเสือ ("옛날 옛적에 호랑이가 있었어요")','ครูมินจุน 80%'),
('기다려져요','เรื่องตอน 2 (จบ) และตอน 3 ("설날이 기다려져요")','ทั้งสองเสียง 80%'),
('죄송합니다','บทเรียน 인사 → ปุ่ม 🐢 ช้า','คลิป "ช้า" ถูกยืดเวลา ×1.5 ตรวจ 40% (ครูมินจุน) — ฟังว่ายังเป็นธรรมชาติไหม'),
('요리하다','บทเรียน → ปุ่ม 🐢 ช้า','คลิป "ช้า" ตรวจ 50–75%'),
('이를 닦다','บทเรียน → ปุ่ม 🐢 ช้า','คลิป "ช้า" ตรวจ 75%'),
('ㄱ ㄴ ㄷ … (ชื่อพยัญชนะ 기역 니은 디귿)','บทเรียน 자음 → ปุ่ม 🔊 / 🐢','พยางค์เดี่ยว/ชื่อตัวอักษรตรวจอัตโนมัติไม่ได้ ต้องฟังด้วยหูเท่านั้น'),
('케이크 / 손풍선 / 곰이요','เรื่อง/สัตว์เลี้ยง','คำที่ผู้พัฒนายังไม่เคยฟังยืนยัน')]
for i,(t,w,why) in enumerate(listen,2):
    for j,v in enumerate([('ครูฮานา และครูมินจุน (เลือกใน ⚙ ตั้งค่า)'),t,w,why],1):
        c=ls.cell(i,j,v); c.font=font(); c.alignment=wrap; c.border=BORDER
    for j in (5,6):
        c=ls.cell(i,j); c.fill=INPUT; c.border=BORDER
dv=DataValidation(type='list',formula1='"✔ ชัดและถูกต้อง,✎ เพี้ยน/อ่านผิด,~ พอฟังได้แต่แปลก"',allow_blank=True); ls.add_data_validation(dv); dv.add(f'E2:E{len(listen)+1}')
ls.freeze_panes='A2'

# ---- summary with formulas
sm=wb.create_sheet('สรุปความคืบหน้า',1)
for j,(h,w) in enumerate([('ชุดข้อความ',22),('จำนวนแถว',12),('ตรวจ ① แล้ว',14),('ตรวจ ② แล้ว',14),('ตรวจ ③ แล้ว',14),('ทำเครื่องหมาย "ควรแก้" (รวม ①②③)',24)],1):
    c=sm.cell(1,j,h); c.font=font(bold=True,color='FFFFFF'); c.fill=HEAD; c.alignment=Alignment(wrap_text=True,vertical='center'); sm.column_dimensions[L(j)].width=w
sm.row_dimensions[1].height=36
nc=len(COLS)
for i,(name,n) in enumerate(summary,2):
    q=f"'{name}'"
    sm.cell(i,1,name).font=font()
    sm.cell(i,2,f"=COUNTA({q}!E2:E{n+1})").font=font()
    for k,off in enumerate((1,3,5)):
        col=L(nc+off); sm.cell(i,3+k,f"=COUNTA({q}!{col}2:{col}{n+1})").font=font()
    sm.cell(i,6,'='+'+'.join(f'COUNTIF({q}!{L(nc+off)}2:{L(nc+off)}{n+1},"✎*")' for off in (1,3,5))).font=font()
t=len(summary)+2
sm.cell(t,1,'รวม').font=font(bold=True)
for j in range(2,7): sm.cell(t,j,f'=SUM({L(j)}2:{L(j)}{t-1})').font=font(bold=True)
sm.cell(t+2,1,'ตัวเลขนี้คำนวณอัตโนมัติจากช่องที่กรอกในแต่ละแท็บ').font=font(italic=True,color='7F7F7F')
wb.save(sys.argv[2])  # เช่น review/hanguk-content-review.xlsx (ต้องมี openpyxl)
print({n:c for n,c in summary})
